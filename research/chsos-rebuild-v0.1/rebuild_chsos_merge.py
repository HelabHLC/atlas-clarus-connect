#!/usr/bin/env python3
"""Deterministic ATLAS RC22 Basis-23 + CHSOS merge rebuild.

This is a research-only spectral mixture search. It preserves every stored
Basis-23 recipe unless a newly computed candidate has a lower CIEDE2000 value.
No output is a physically measured or production-approved paint recipe.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize
from scipy.spatial import cKDTree


METHOD = "KS_PROXY_D50_2DEG_400_700_V2_REBUILD"
MASTER_SHA256 = "8283ab91b10f89ac758d09ecf5fb4d6343536600a06dd468b1cc1ecf4ec747c4"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_js_assignment(path: Path, prefix: str) -> dict:
    text = path.read_text(encoding="utf-8")
    if not text.startswith(prefix):
        raise ValueError(f"Unexpected JavaScript prefix in {path}")
    end = text.rfind(";")
    if end < len(prefix):
        raise ValueError(f"Missing assignment terminator in {path}")
    return json.loads(text[len(prefix):end])


def km_from_reflectance(r: np.ndarray) -> np.ndarray:
    r = np.clip(np.asarray(r, dtype=np.float64), 1e-6, 1.0)
    return (1.0 - r) ** 2 / (2.0 * r)


def reflectance_from_km(ks: np.ndarray) -> np.ndarray:
    ks = np.maximum(np.asarray(ks, dtype=np.float64), 0.0)
    return np.clip(1.0 + ks - np.sqrt(ks * ks + 2.0 * ks), 0.0, 1.0)


def f_lab(t: np.ndarray) -> np.ndarray:
    delta = 6.0 / 29.0
    return np.where(t > delta**3, np.cbrt(t), t / (3 * delta**2) + 4.0 / 29.0)


@dataclass(frozen=True)
class Colorimetry:
    wx: np.ndarray
    wy: np.ndarray
    wz: np.ndarray
    white: np.ndarray

    def lab(self, reflectance: np.ndarray) -> np.ndarray:
        r = np.asarray(reflectance, dtype=np.float64)
        xyz = np.stack((r @ self.wx, r @ self.wy, r @ self.wz), axis=-1)
        q = f_lab(xyz / self.white)
        return np.stack((116 * q[..., 1] - 16,
                         500 * (q[..., 0] - q[..., 1]),
                         200 * (q[..., 1] - q[..., 2])), axis=-1)


def delta_e_2000(lab1: np.ndarray, lab2: np.ndarray) -> np.ndarray:
    """Vectorized Sharma et al. CIEDE2000, kL=kC=kH=1."""
    a = np.asarray(lab1, dtype=np.float64)
    b = np.asarray(lab2, dtype=np.float64)
    l1, a1, b1 = np.moveaxis(a, -1, 0)
    l2, a2, b2 = np.moveaxis(b, -1, 0)
    c1 = np.hypot(a1, b1); c2 = np.hypot(a2, b2); cb = (c1 + c2) / 2
    g = 0.5 * (1 - np.sqrt(cb**7 / (cb**7 + 25**7)))
    ap1 = (1 + g) * a1; ap2 = (1 + g) * a2
    cp1 = np.hypot(ap1, b1); cp2 = np.hypot(ap2, b2)
    hp1 = np.mod(np.degrees(np.arctan2(b1, ap1)), 360)
    hp2 = np.mod(np.degrees(np.arctan2(b2, ap2)), 360)
    hp1 = np.where((ap1 == 0) & (b1 == 0), 0, hp1)
    hp2 = np.where((ap2 == 0) & (b2 == 0), 0, hp2)
    dl = l2 - l1; dc = cp2 - cp1; dh_raw = hp2 - hp1
    dh = np.where(cp1 * cp2 == 0, 0,
         np.where(np.abs(dh_raw) <= 180, dh_raw,
         np.where(dh_raw > 180, dh_raw - 360, dh_raw + 360)))
    dH = 2 * np.sqrt(cp1 * cp2) * np.sin(np.radians(dh / 2))
    lb = (l1 + l2) / 2; cpb = (cp1 + cp2) / 2
    hsum = hp1 + hp2
    hb = np.where(cp1 * cp2 == 0, hsum,
         np.where(np.abs(hp1 - hp2) <= 180, hsum / 2,
         np.where(hsum < 360, (hsum + 360) / 2, (hsum - 360) / 2)))
    t = (1 - 0.17*np.cos(np.radians(hb-30)) + 0.24*np.cos(np.radians(2*hb))
         + 0.32*np.cos(np.radians(3*hb+6)) - 0.20*np.cos(np.radians(4*hb-63)))
    sl = 1 + 0.015*(lb-50)**2 / np.sqrt(20+(lb-50)**2)
    sc = 1 + 0.045*cpb; sh = 1 + 0.015*cpb*t
    rt = -2*np.sqrt(cpb**7/(cpb**7+25**7))*np.sin(np.radians(
         60*np.exp(-((hb-275)/25)**2)))
    x, y, z = dl/sl, dc/sc, dH/sh
    return np.sqrt(np.maximum(0, x*x + y*y + z*z + rt*y*z))


def slug_id(name: str) -> str:
    s = name.upper().replace("–", "-")
    s = re.sub(r"[^A-Z0-9]+", "_", s).strip("_")
    return "CHSOS_GORGIAS_" + s


def load_chsos_xlsx(path: Path) -> tuple[list[dict], Colorimetry]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Spektren_1nm"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[9])
    wavelengths = np.array([r[0] for r in rows[10:] if r[0] is not None], dtype=float)
    mask = (wavelengths >= 400) & (wavelengths <= 700) & ((wavelengths % 10) == 0)
    # The workbook embeds CIE D50/1931 2-degree values. Ten-nm trapezoidal
    # weights are re-normalized on the declared 400-700 nm model interval.
    d50 = np.array([r[1] for r in rows[10:]], dtype=float)[mask]
    xbar = np.array([r[2] for r in rows[10:]], dtype=float)[mask]
    ybar = np.array([r[3] for r in rows[10:]], dtype=float)[mask]
    zbar = np.array([r[4] for r in rows[10:]], dtype=float)[mask]
    trap = np.ones(mask.sum()); trap[[0, -1]] = 0.5
    norm = np.sum(d50 * ybar * trap)
    wx = 100*d50*xbar*trap/norm; wy = 100*d50*ybar*trap/norm; wz = 100*d50*zbar*trap/norm
    colorimetry = Colorimetry(wx, wy, wz, np.array([wx.sum(), 100.0, wz.sum()]))
    samples = []
    for col in range(10, len(header)):
        name = header[col]
        if not name or str(name).strip().lower() == "acrylic binder":
            continue
        refl = np.array([r[col] for r in rows[10:]], dtype=float)[mask]
        if len(refl) != 31 or not np.isfinite(refl).all():
            raise ValueError(f"Invalid spectrum for {name}")
        samples.append({"basis_id": slug_id(str(name)), "sample_title": str(name),
                        "source_family": "CHSOS_GORGIAS_FORS",
                        "reflectance_400_700": refl.tolist()})
    if len(samples) != 65:
        raise ValueError(f"Expected 65 CHSOS colorants, got {len(samples)}")
    return samples, colorimetry


def make_candidate_pool(n_basis: int, old_count: int, random_count: int,
                        seed: int) -> tuple[np.ndarray, np.ndarray]:
    """Return deterministic sparse supports and weights, requiring new CHSOS."""
    rng = np.random.default_rng(seed)
    supports: list[np.ndarray] = []
    weights: list[np.ndarray] = []
    # Every new 1C and every pair with a new member on a 10% grid.
    for i in range(old_count, n_basis):
        supports.append(np.array([i, -1, -1, -1])); weights.append(np.array([1., 0, 0, 0]))
    for i in range(n_basis):
        for j in range(max(i+1, old_count), n_basis):
            if i == j: continue
            for pct in range(10, 100, 10):
                supports.append(np.array([i, j, -1, -1]))
                weights.append(np.array([pct/100, 1-pct/100, 0, 0]))
    # Fixed-seed 3C/4C Dirichlet candidates. At least one component is new.
    for _ in range(random_count):
        k = 3 if rng.random() < 0.35 else 4
        idx = rng.choice(n_basis, k, replace=False)
        if np.all(idx < old_count):
            idx[rng.integers(k)] = rng.integers(old_count, n_basis)
            if len(set(idx.tolist())) != k:
                continue
        w = rng.dirichlet(np.ones(k))
        s = np.full(4, -1, dtype=np.int32); q = np.zeros(4)
        s[:k] = idx; q[:k] = w
        supports.append(s); weights.append(q)
    return np.vstack(supports).astype(np.int32), np.vstack(weights)


def pool_labs(supports: np.ndarray, weights: np.ndarray, ks: np.ndarray,
              cm: Colorimetry, chunk: int = 50000) -> np.ndarray:
    out = np.empty((len(supports), 3), dtype=np.float64)
    for start in range(0, len(supports), chunk):
        stop = min(start + chunk, len(supports)); s = supports[start:stop]; w = weights[start:stop]
        mixed = np.zeros((stop-start, ks.shape[1]))
        for c in range(4):
            valid = s[:, c] >= 0
            mixed[valid] += w[valid, c, None] * ks[s[valid, c]]
        out[start:stop] = cm.lab(reflectance_from_km(mixed))
    return out


def refine(support: np.ndarray, initial: np.ndarray, ks: np.ndarray,
           target: np.ndarray, cm: Colorimetry) -> tuple[np.ndarray, float, np.ndarray]:
    valid = support >= 0; idx = support[valid]; x0 = initial[valid]
    x0 = np.maximum(x0, 1e-6); x0 /= x0.sum()
    def evaluate(x):
        refl = reflectance_from_km(x @ ks[idx])
        lab = cm.lab(refl)
        return float(delta_e_2000(lab, target)), refl
    res = minimize(lambda x: evaluate(x)[0], x0, method="SLSQP",
                   bounds=[(0, 1)]*len(idx), constraints={"type":"eq", "fun":lambda x:x.sum()-1},
                   options={"maxiter":100, "ftol":1e-10, "disp":False})
    x = np.maximum(res.x if res.success else x0, 0); x /= x.sum()
    de, refl = evaluate(x)
    # RC22's documented fine grid has a 0.5 percentage-point floor. Removing
    # smaller numerical residues prevents an old-basis solution from being
    # mislabeled as a CHSOS-derived improvement.
    keep = x >= 0.005
    if not keep.all():
        idx = idx[keep]; x = x[keep]; x /= x.sum(); de, refl = evaluate_reduced(idx, x, ks, target, cm)
    return np.column_stack((idx, x)), de, refl


def evaluate_reduced(idx, x, ks, target, cm):
    refl = reflectance_from_km(x @ ks[idx]); lab = cm.lab(refl)
    return float(delta_e_2000(lab, target)), refl


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--rc22-data", type=Path, required=True)
    ap.add_argument("--atlas-data", type=Path, required=True)
    ap.add_argument("--basis23", type=Path, required=True)
    ap.add_argument("--chsos-xlsx", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--random-candidates", type=int, default=250000)
    ap.add_argument("--query-k", type=int, default=6)
    ap.add_argument("--seed", type=int, default=20260915)
    ap.add_argument("--limit", type=int, default=0, help="Development limit; 0 means all rows")
    ap.add_argument("--start", type=int, default=0, help="First atlas_row_id, inclusive")
    ap.add_argument("--end", type=int, default=0, help="Last atlas_row_id, exclusive; 0 means all")
    args = ap.parse_args()

    old = read_js_assignment(args.rc22_data, "window.ATLAS_BASIS23_DATA=")
    atlas = read_js_assignment(args.atlas_data, "window.ATLAS_CLARUS_DATA=")
    if atlas["master_sha256"] != MASTER_SHA256 or len(atlas["colors"]) != 13283:
        raise ValueError("ATLAS master identity mismatch")
    old_basis = json.loads(args.basis23.read_text(encoding="utf-8"))
    new_basis, cm = load_chsos_xlsx(args.chsos_xlsx)
    basis = old_basis + new_basis
    ks = km_from_reflectance(np.array([b["reflectance_400_700"] for b in basis]))
    supports, weights = make_candidate_pool(len(basis), len(old_basis), args.random_candidates, args.seed)
    labs = pool_labs(supports, weights, ks, cm)
    tree = cKDTree(labs)
    targets = np.array([c["lab"] for c in atlas["colors"]], dtype=float)
    start = args.start
    end = args.end or len(targets)
    if args.limit:
        end = min(end, start + args.limit)
    if not (0 <= start < end <= len(targets)):
        raise ValueError(f"Invalid row range {start}:{end}")
    count = end - start
    _, near = tree.query(targets[start:end], k=min(args.query_k, len(labs)))
    if near.ndim == 1: near = near[:, None]

    basis_by_id = {b["basis_id"]: b for b in basis}
    results = []
    improved = rescued = 0
    for local_id, row_id in enumerate(range(start, end)):
        prior = old["rows"][row_id]
        best_de = float(prior["de00"]); best_parts = None; best_refl = None
        seen = set()
        for pi in near[local_id]:
            key = tuple(sorted(int(x) for x in supports[pi] if x >= 0))
            if key in seen: continue
            seen.add(key)
            parts, de, refl = refine(supports[pi], weights[pi], ks, targets[row_id], cm)
            if not np.any(parts[:, 0] >= len(old_basis)):
                continue
            if de + 1e-9 < best_de:
                best_de, best_parts, best_refl = de, parts, refl
        if best_parts is None:
            results.append({**prior, "merge_decision":"RETAINED_BASIS23"})
            continue
        improved += 1
        if float(prior["de00"]) > 5 and best_de <= 5: rescued += 1
        components = []
        for idx, fraction in best_parts:
            b = basis[int(idx)]
            components.append({"basis_id":b["basis_id"], "name":b["sample_title"],
                "percent":round(float(fraction)*100, 6), "source_family":b["source_family"],
                "source_url":"https://chsopensource.org/products/pigments-checker/" if b["source_family"].startswith("CHSOS") else "https://github.com/miciwan/PaintMixing"})
        results.append({
            "source_atlas_row_id":row_id, "reference":atlas["colors"][row_id]["ref"],
            "basis_version":"ATLAS_RC22_BASIS23_CHSOS_GORGIAS_v0_1",
            "component_count":len(components), "components":components,
            "de00":round(best_de, 9), "previous_basis23_de00":float(prior["de00"]),
            "computational_tolerance":5.0,
            "computational_tolerance_status":"WITHIN_COMPUTATIONAL_TOLERANCE" if best_de <= 5 else "OUTSIDE_COMPUTATIONAL_TOLERANCE",
            "search_optimality":"BEST_FOUND_DETERMINISTIC_HEURISTIC_NOT_GLOBAL_PROOF",
            "measured_qc_status":"NOT_MEASURED", "production_approval":"NOT_SUPPORTED",
            "merge_decision":"IMPROVED_WITH_CHSOS_GORGIAS"
        })
        if (local_id + 1) % 250 == 0:
            print(f"{local_id+1}/{count} rows ({start}:{end}); improved={improved}; rescued={rescued}", file=sys.stderr)

    within = sum(float(r["de00"]) <= 5 for r in results)
    payload = {
        "registry": {
            "dataset":"ATLAS_RC22_Basis23_CHSOS_Merged_Recipes_v0_1",
            "created_utc":datetime.now(timezone.utc).isoformat(),
            "atlas_master_sha256":MASTER_SHA256,
            "method":METHOD, "seed":args.seed,
            "random_candidate_count":args.random_candidates, "query_k":args.query_k,
            "row_range":[start, end], "rows":len(results), "within_de00_5":within,
            "outside_de00_5":len(results)-within, "improved_rows":improved,
            "rescued_to_le5":rescued, "regressions":0,
            "status":"MODEL_ONLY_NOT_PHYSICALLY_VALIDATED",
            "input_sha256": {"rc22_data":sha256(args.rc22_data), "atlas_data":sha256(args.atlas_data),
                             "basis23":sha256(args.basis23), "chsos_xlsx":sha256(args.chsos_xlsx)}
        },
        "basis_registry":basis,
        "rows":results
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",",":")), encoding="utf-8")
    print(json.dumps(payload["registry"], indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
